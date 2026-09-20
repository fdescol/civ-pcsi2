#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sync.py -- Synchronise le colloscope depuis Google Sheets.

Usage :
    python sync.py          # telecharge, compare, demande confirmation
    python sync.py --auto   # met a jour sans confirmation (script batch)
    python sync.py --check  # compare uniquement, sans proposer la mise a jour
"""
import argparse
import shutil
import subprocess
import sys
import urllib.request
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
SPREADSHEET_ID = "1eL0ZbPJXE0zdFNRiFXLSVzSLKFqk2T9Y7VNR4qHAITM"
EXPORT_URL = (
    "https://docs.google.com/spreadsheets/d/"
    + SPREADSHEET_ID
    + "/export?format=xlsx"
)
LOCAL_XLSX  = Path("Colles TDs et TPs.xlsx")
TMP_XLSX    = Path("_downloaded_tmp.xlsx")
CONVERT_PY  = Path("convert.py")

# Seules les feuilles lues par convert.py sont comparees
SHEETS_TO_COMPARE = ["Colloscope", "Trinomes", "Semaines", "TPTD"]
# Noms reels dans le fichier (avec accents) pour la recherche
SHEETS_ACCENTED   = ["Colloscope", "Trinômes", "Semaines", "TPTD"]


# ---------------------------------------------------------------------------
def download_sheet() -> Path:
    """Telecharge le Google Sheets public et l'ecrit dans TMP_XLSX."""
    print("Telechargement depuis Google Sheets...", end=" ", flush=True)
    req = urllib.request.Request(
        EXPORT_URL, headers={"User-Agent": "Mozilla/5.0"}
    )
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = resp.read()
    TMP_XLSX.write_bytes(data)
    print(f"OK ({len(data) // 1024} Ko)")
    return TMP_XLSX


# ---------------------------------------------------------------------------
def _cell_addr(row: int, col: int) -> str:
    """Convertit (row, col) en adresse Excel style 'A1'."""
    result = ""
    while col > 0:
        col, remainder = divmod(col - 1, 26)
        result = chr(65 + remainder) + result
    return f"{result}{row}"


def compare_workbooks(local_path: Path, remote_path: Path) -> list:
    """
    Compare cellule par cellule les feuilles utiles des deux classeurs.
    Retourne une liste de dicts {sheet, cell, local_val, remote_val}.
    """
    wb_local  = openpyxl.load_workbook(local_path,  data_only=True)
    wb_remote = openpyxl.load_workbook(remote_path, data_only=True)

    diffs = []
    for sheet_name in SHEETS_ACCENTED:
        if sheet_name not in wb_local.sheetnames:
            print(f"  [?] Feuille '{sheet_name}' absente du fichier local — ignoree")
            continue
        if sheet_name not in wb_remote.sheetnames:
            print(f"  [?] Feuille '{sheet_name}' absente du fichier distant — ignoree")
            continue

        ws_l = wb_local[sheet_name]
        ws_r = wb_remote[sheet_name]

        max_row = max(ws_l.max_row or 1, ws_r.max_row or 1)
        max_col = max(ws_l.max_column or 1, ws_r.max_column or 1)

        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                v_local  = ws_l.cell(row=r, column=c).value
                v_remote = ws_r.cell(row=r, column=c).value
                # Normaliser les espaces et None
                v_local  = str(v_local).strip()  if v_local  is not None else ""
                v_remote = str(v_remote).strip() if v_remote is not None else ""
                if v_local != v_remote:
                    diffs.append({
                        "sheet":      sheet_name,
                        "cell":       _cell_addr(r, c),
                        "local_val":  v_local,
                        "remote_val": v_remote,
                    })
    return diffs


# ---------------------------------------------------------------------------
def print_report(diffs: list) -> None:
    """Affiche le rapport de differences groupe par feuille."""
    if not diffs:
        print("Aucune difference detectee sur les feuilles comparees.")
        return

    from itertools import groupby
    key = lambda d: d["sheet"]
    diffs_sorted = sorted(diffs, key=key)

    total = 0
    for sheet, group in groupby(diffs_sorted, key=key):
        items = list(group)
        total += len(items)
        print(f"\n  Feuille '{sheet}' : {len(items)} difference(s)")
        # Limiter l'affichage a 20 diffs par feuille pour rester lisible
        shown = items[:20]
        for d in shown:
            lv = d["local_val"]  if d["local_val"]  != "" else "(vide)"
            rv = d["remote_val"] if d["remote_val"] != "" else "(vide)"
            print(f"    {d['cell']:6s}  local={lv!r:25s}  distant={rv!r}")
        if len(items) > 20:
            print(f"    ... et {len(items) - 20} autres differences non affichees")

    print(f"\n  TOTAL : {total} cellule(s) differente(s) sur {len(set(d['sheet'] for d in diffs))} feuille(s)")


# ---------------------------------------------------------------------------
def apply_update() -> None:
    """Remplace le fichier local par le telechargement et relance convert.py."""
    shutil.copy(TMP_XLSX, LOCAL_XLSX)
    print(f"Fichier local mis a jour : {LOCAL_XLSX}")
    print("Relancement de convert.py...")
    result = subprocess.run([sys.executable, str(CONVERT_PY)], check=True)
    if result.returncode == 0:
        print("\ndata.json regenere avec succes.")
        print("Prochaines etapes :")
        print('  git add data.json "Colles TDs et TPs.xlsx"')
        print('  git commit -m "Update planning data"')
        print("  git push")


# ---------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(
        description="Synchronise le colloscope depuis Google Sheets."
    )
    parser.add_argument(
        "--auto",
        action="store_true",
        help="Met a jour sans demander de confirmation",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="Compare uniquement, sans proposer la mise a jour",
    )
    args = parser.parse_args()

    if not LOCAL_XLSX.exists():
        print(f"ERREUR : fichier local introuvable : {LOCAL_XLSX}")
        sys.exit(1)
    if not CONVERT_PY.exists() and not args.check:
        print(f"ERREUR : convert.py introuvable dans le repertoire courant")
        sys.exit(1)

    try:
        download_sheet()

        print("Comparaison des feuilles...", end=" ", flush=True)
        diffs = compare_workbooks(LOCAL_XLSX, TMP_XLSX)
        print("OK")

        print_report(diffs)

        if not diffs:
            sys.exit(0)

        if args.check:
            sys.exit(1)  # differences trouvees, mais mode check seul

        if args.auto:
            apply_update()
        else:
            answer = input("\nMettre a jour le fichier local et relancer convert.py ? [o/N] ").strip().lower()
            if answer in ("o", "oui", "y", "yes"):
                apply_update()
            else:
                print("Mise a jour annulee. Le fichier local n'a pas ete modifie.")

    except KeyboardInterrupt:
        print("\nInterrompu.")
        sys.exit(1)
    except urllib.error.URLError as e:
        print(f"\nERREUR reseau : {e}")
        sys.exit(1)
    except subprocess.CalledProcessError:
        print("\nERREUR : convert.py a echoue. Verifiez le fichier xlsx.")
        sys.exit(1)
    finally:
        if TMP_XLSX.exists():
            TMP_XLSX.unlink()


if __name__ == "__main__":
    main()
